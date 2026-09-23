"""
Blackboard Ultra quiz question harvester.

Repeatedly starts a self-test attempt, records every question shown, submits
the attempt (typing '0' into the last text box), and collects every unique
question into a Word document.

Per attempt the flow is:
    "Start attempt N" -> scroll to bottom -> type '0' into last text box
    -> "Submit" -> "Submit" (confirm dialog) -> "Close" -> repeat

Usage:
    pip install playwright python-docx
    playwright install chromium
    python scrape_quiz.py                       # log in, open the quiz page, press Enter
    python scrape_quiz.py --url "<details page URL>" --max-attempts 50

Progress is saved to the output folder after every attempt, so you can stop
(Ctrl+C) and re-run later; already-seen questions are never added twice.
"""

VERSION = "4"

import argparse
import hashlib
import json
import re
import sys
import time
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from playwright.sync_api import Error as PlaywrightError
from playwright.sync_api import TimeoutError as PlaywrightTimeout
from playwright.sync_api import sync_playwright

# JavaScript run inside the attempt page. It finds every answer control, groups
# controls belonging to the same question (radio buttons by name, checkboxes by
# their enclosing fieldset/group), then climbs from each group to the largest
# ancestor that contains no other question's controls. That ancestor is the
# question block (text, images, answer box). Each block is tagged with a
# data attribute so Python can screenshot it afterwards.
FIND_QUESTIONS_JS = r"""
() => {
  const visible = el => {
    const r = el.getBoundingClientRect();
    const s = getComputedStyle(el);
    return (r.width > 0 || r.height > 0) && s.visibility !== 'hidden' && s.display !== 'none';
  };
  const controls = [...document.querySelectorAll(
    'input[type=text], input[type=number], input:not([type]), textarea, select, ' +
    'input[type=radio], input[type=checkbox], [contenteditable=true], [role=textbox]'
  )].filter(el => !el.closest('header, nav, footer, [role=dialog]'))
    // hidden radios/checkboxes are often styled replacements - keep them
    .filter(el => visible(el) || el.type === 'radio' || el.type === 'checkbox');

  const groupKey = el => {
    if (el.type === 'radio' && el.name) return 'radio:' + el.name;
    if (el.type === 'radio' || el.type === 'checkbox') {
      const g = el.closest('fieldset, [role=group], [role=radiogroup], ul, ol');
      if (g) {
        if (!g.dataset.qsGroup) g.dataset.qsGroup = Math.random().toString(36).slice(2);
        return 'grp:' + g.dataset.qsGroup;
      }
    }
    if (!el.dataset.qsGroup) el.dataset.qsGroup = Math.random().toString(36).slice(2);
    return 'el:' + el.dataset.qsGroup;
  };

  const groups = new Map();
  for (const el of controls) {
    // skip inputs nested inside a rich-text editor that is itself a control
    if (el.parentElement && el.parentElement.closest('[contenteditable=true]')) continue;
    const k = groupKey(el);
    if (!groups.has(k)) groups.set(k, []);
    groups.get(k).push(el);
  }

  const keyOf = new Map();
  for (const [k, els] of groups) for (const el of els) keyOf.set(el, k);
  const foreignInside = (node, k) => {
    for (const [el, kk] of keyOf) if (kk !== k && node.contains(el)) return true;
    return false;
  };

  const blocks = [];
  for (const [k, els] of groups) {
    // start from the lowest common ancestor of all controls in the group
    let node = els[0];
    while (!els.every(e => node.contains(e))) node = node.parentElement;
    if (groups.size > 1) {
      while (node.parentElement && node.parentElement !== document.body &&
             !foreignInside(node.parentElement, k)) {
        node = node.parentElement;
      }
    } else {
      // only one question on the page: climb a few levels to pick up its text
      for (let i = 0; i < 6 && node.parentElement && node.parentElement !== document.body; i++) {
        node = node.parentElement;
        if ((node.innerText || '').trim().length > 40) break;
      }
    }
    if (!blocks.includes(node)) blocks.push(node);
  }

  blocks.sort((a, b) => a.compareDocumentPosition(b) & Node.DOCUMENT_POSITION_FOLLOWING ? -1 : 1);

  return blocks.map((node, i) => {
    node.setAttribute('data-qs-block', String(i));
    const imgs = [...node.querySelectorAll('img')].map(img => {
      const src = img.currentSrc || img.src || '';
      return src.startsWith('data:') ? src.slice(0, 200) : src.split('?')[0].split('#')[0];
    });
    return { index: i, text: node.innerText || '', images: imgs };
  });
}
"""

SCROLL_BOTTOM_JS = r"""
() => {
  window.scrollTo(0, document.body.scrollHeight);
  for (const el of document.querySelectorAll('*')) {
    const s = getComputedStyle(el);
    if ((s.overflowY === 'auto' || s.overflowY === 'scroll') && el.scrollHeight > el.clientHeight) {
      el.scrollTop = el.scrollHeight;
    }
  }
}
"""

# Lines that are UI chrome rather than question content.
NOISE_LINES = re.compile(
    r"^(\d+\s+of\s+\d+\s+questions?\s+remaining|question\s+\d+(\s+of\s+\d+)?|"
    r"\d+(\.\d+)?\s+points?|flag( question)?|answer|your answer|type your answer here.*)$",
    re.IGNORECASE,
)


def clean_text(text):
    lines = [ln.strip() for ln in text.splitlines()]
    lines = [ln for ln in lines if ln and not NOISE_LINES.match(ln)]
    return "\n".join(lines)


NUMBER = re.compile(r"[-+]?\d+(?:[.,]\d+)*(?:[eE][-+]?\d+)?")


def question_key(text, images, ignore_numbers=True):
    norm = re.sub(r"\s+", " ", text).strip().lower()
    if ignore_numbers:
        # questions whose numbers are randomised each attempt count as the same question
        norm = NUMBER.sub("#", norm)
    return hashlib.sha1((norm + "|" + "|".join(sorted(images))).encode("utf-8")).hexdigest()


class Store:
    """Keeps seen questions on disk (JSON + screenshots) and rebuilds the .docx."""

    def __init__(self, out_dir, title, ignore_numbers=True):
        self.dir = Path(out_dir)
        self.img_dir = self.dir / "screenshots"
        self.img_dir.mkdir(parents=True, exist_ok=True)
        self.json_path = self.dir / "questions.json"
        self.docx_path = self.dir / "questions.docx"
        self.title = title
        self.questions = []
        if self.json_path.exists():
            self.questions = json.loads(self.json_path.read_text(encoding="utf-8"))
        self.ignore_numbers = ignore_numbers
        # recompute keys so saved progress follows the current duplicate rule
        self.keys = set()
        unique = []
        for q in self.questions:
            q["key"] = question_key(q["text"], q["images"], ignore_numbers)
            if q["key"] not in self.keys:
                self.keys.add(q["key"])
                unique.append(q)
        self.questions = unique

    def has(self, key):
        return key in self.keys

    def add(self, key, text, images, screenshot, attempt):
        self.questions.append({
            "key": key, "text": text, "images": images,
            "screenshot": screenshot, "first_seen_attempt": attempt,
        })
        self.keys.add(key)

    def save(self):
        self.json_path.write_text(json.dumps(self.questions, indent=2), encoding="utf-8")
        doc = Document()
        doc.styles["Normal"].font.size = Pt(11)
        doc.add_heading(self.title, level=0)
        doc.add_paragraph(f"{len(self.questions)} unique questions")
        for n, q in enumerate(self.questions, 1):
            doc.add_heading(f"Question {n}", level=2)
            for para in q["text"].split("\n"):
                doc.add_paragraph(para)
            shot = q.get("screenshot")
            if shot and (self.dir / shot).exists():
                doc.add_picture(str(self.dir / shot), width=Inches(6))
        tmp = self.docx_path.with_suffix(".tmp.docx")
        doc.save(tmp)
        try:
            tmp.replace(self.docx_path)
        except PermissionError:
            # file is open in Word on Windows - keep the temp copy instead
            print(f"  ! {self.docx_path.name} is open elsewhere; saved to {tmp.name}")


def all_frames(page):
    return [page.main_frame] + [f for f in page.frames if f != page.main_frame]


def find_visible(page, make_locators, timeout=15.0):
    """Poll all frames until one of the locators built by make_locators(frame) is visible."""
    deadline = time.time() + timeout
    while time.time() < deadline:
        for frame in all_frames(page):
            try:
                for loc in make_locators(frame):
                    for i in range(loc.count()):
                        item = loc.nth(i)
                        if item.is_visible():
                            return item
            except PlaywrightError:
                pass  # frame navigated/detached mid-check
        time.sleep(0.25)
    return None


def retry(action, what, tries=5):
    """Run action(try_no), retrying when Blackboard re-renders the element mid-action."""
    last = None
    for n in range(tries):
        try:
            return action(n)
        except PlaywrightError as e:
            last = e
            time.sleep(1.0)
    raise RuntimeError(f"{what} failed after {tries} tries: {str(last).splitlines()[0]}")


def click(page, make_locators, what, timeout=15.0):
    def do(n):
        loc = find_visible(page, make_locators, timeout)
        if loc is None:
            raise RuntimeError(f"Could not find {what}")
        # Playwright scrolls into view itself; on the last tries skip its stability checks
        loc.click(timeout=5000, force=n >= 3)
    retry(do, f"Clicking {what}")


def attempt_frame(page, timeout=30.0):
    """Return the frame that holds the quiz's answer controls."""
    deadline = time.time() + timeout
    while time.time() < deadline:
        for frame in all_frames(page):
            try:
                if frame.locator("input[type=text], input:not([type]), textarea, input[type=radio], "
                                 "input[type=checkbox], [contenteditable=true]").count():
                    if frame.get_by_role("button", name=re.compile(r"^\s*submit\s*$", re.I)).count():
                        return frame
            except PlaywrightError:
                pass
        time.sleep(0.5)
    return None


def load_all_questions(frame):
    """Scroll to the bottom until the number of answer controls stops growing (lazy loading)."""
    last = -1
    for _ in range(20):
        frame.evaluate(SCROLL_BOTTOM_JS)
        time.sleep(0.6)
        count = frame.locator("input, textarea, [contenteditable=true]").count()
        if count == last:
            break
        last = count


START_RE = re.compile(r"(start|continue|resume) attempt", re.I)


def run_attempt(page, details_url, store, attempt_no, args):
    # 1. Make sure we are on the details page and press "Start attempt N".
    start = find_visible(page, start_button, 5)
    if start is None:
        page.goto(details_url, wait_until="domcontentloaded")
        start = find_visible(page, start_button, 30)
        if start is None:
            raise RuntimeError("Could not find the 'Start attempt' button")
    click(page, lambda f: [f.get_by_role(role, name=START_RE) for role in ("button", "link")],
          "the Start attempt button")

    # Some quizzes show an extra "Start attempt"/"Begin" confirmation dialog.
    confirm_btn = lambda f: [f.locator("[role=dialog]").get_by_role(
        "button", name=re.compile(r"^(start attempt|start|begin|continue)$", re.I))]
    if find_visible(page, confirm_btn, 3) is not None:
        click(page, confirm_btn, "the Start confirmation button", timeout=3)

    # 2. Wait for the attempt to load and scroll to the bottom.
    frame = attempt_frame(page)
    if frame is None:
        raise RuntimeError("Attempt page did not load (no answer boxes / Submit button found)")
    time.sleep(args.settle)
    load_all_questions(frame)

    # 3. Harvest questions.
    blocks = frame.evaluate(FIND_QUESTIONS_JS)
    new = 0
    for b in blocks:
        text = clean_text(b["text"])
        if len(text) < 5 and not b["images"]:
            continue
        key = question_key(text, b["images"], store.ignore_numbers)
        if store.has(key):
            continue
        shot = None
        if not args.no_screenshots:
            shot = f"screenshots/q{len(store.questions) + 1:04d}.png"
            el = frame.locator(f"[data-qs-block='{b['index']}']")
            try:
                retry(lambda n: el.screenshot(path=str(store.dir / shot), timeout=10000,
                                              animations="disabled"), "Screenshot", tries=3)
            except RuntimeError as e:
                print(f"  ! {e} - saving the question as text only")
                shot = None
        store.add(key, text, b["images"], shot, attempt_no)
        new += 1

    # 4. Type '0' into the last text box.
    frame.evaluate(SCROLL_BOTTOM_JS)
    time.sleep(1.0)  # let the scroll finish so elements stop moving

    def type_zero(n):
        # re-find the box on every try: Blackboard replaces elements while it re-renders
        boxes = frame.locator("input[type=text], input[type=number], input:not([type]), textarea, "
                              "[contenteditable=true]")
        visible = [i for i in range(boxes.count()) if boxes.nth(i).is_visible()]
        if not visible:
            return
        box = boxes.nth(visible[-1])
        box.fill("0", timeout=5000, force=n >= 3)
        box.press("Tab")  # blur so Blackboard autosaves the answer

    retry(type_zero, "Typing 0 into the last text box")
    time.sleep(1.0)

    # 5. Submit -> confirm Submit -> Close.
    click(page, lambda f: [f.get_by_role("button", name=re.compile(r"^\s*submit\s*$", re.I))],
          "the Submit button")
    click(page, lambda f: [f.locator("[role=dialog], [role=alertdialog], .modal").get_by_role(
        "button", name=re.compile(r"^\s*submit\s*$", re.I))], "the Submit confirmation button")
    click(page, lambda f: [f.locator("[role=dialog], [role=alertdialog], .modal").get_by_role(
        "button", name=re.compile(r"^\s*close\s*$", re.I))], "the Close button", timeout=30)
    time.sleep(args.settle)
    return len(blocks), new


def current_page(ctx, page):
    """The tab to drive: the given one if still open, else the newest open tab."""
    if page is not None and not page.is_closed():
        return page
    open_pages = [p for p in ctx.pages if not p.is_closed()]
    return open_pages[-1] if open_pages else ctx.new_page()


def start_button(f):
    return [f.get_by_role(r, name=START_RE) for r in ("button", "link")]


def wait_for_user_on_quiz_page(ctx):
    """Ask the user to open the quiz page, then find whichever tab it is in."""
    while True:
        print("\nLog in if needed and open the quiz's details page (the one with 'Start attempt').")
        input("Press Enter here when that page is showing... ")
        for p in reversed(ctx.pages):
            if not p.is_closed() and find_visible(p, start_button, 2) is not None:
                p.bring_to_front()
                return p, p.url
        print("Couldn't see a 'Start attempt' button in any tab - please try again.")


def back_to_quiz(ctx, page, details_url):
    """Get back to the quiz details page. Returns the page, or None if that failed."""
    for p in reversed(ctx.pages):  # maybe it's already showing in some tab
        if not p.is_closed() and find_visible(p, start_button, 1) is not None:
            return p
    page = current_page(ctx, page)
    try:
        page.goto(details_url, wait_until="domcontentloaded")
        if find_visible(page, start_button, 30) is not None:
            return page
    except PlaywrightError:
        pass
    print("  Could not get back to the quiz page automatically.")
    return None


def main():
    ap = argparse.ArgumentParser(description="Harvest Blackboard quiz questions into a Word document.")
    ap.add_argument("--url", help="URL of the quiz's 'Assessment Details' page (the one with 'Start attempt'). "
                                  "If omitted you navigate there yourself and press Enter.")
    ap.add_argument("--out", default="quiz_output", help="Output folder (default: quiz_output)")
    ap.add_argument("--title", default="Quiz Question Bank", help="Title at the top of the Word document")
    ap.add_argument("--max-attempts", type=int, default=100, help="Stop after this many attempts (default 100)")
    ap.add_argument("--stop-after", type=int, default=15,
                    help="Stop after this many attempts in a row with no new questions (default 15)")
    ap.add_argument("--settle", type=float, default=2.0, help="Seconds to wait for pages to settle (default 2)")
    ap.add_argument("--keep-number-variants", action="store_true",
                    help="Save a question again when only its numbers differ (default: ignore number changes)")
    ap.add_argument("--no-screenshots", action="store_true", help="Text only, no question screenshots")
    ap.add_argument("--browser", default=None, choices=["msedge", "chrome"],
                    help="Use an installed Edge/Chrome instead of Playwright's Chromium")
    ap.add_argument("--executable", default=None, help="Path to a Chromium/Chrome/Edge executable to use")
    ap.add_argument("--profile", default=".bb_browser_profile",
                    help="Browser profile folder, so your login is remembered between runs")
    args = ap.parse_args()

    print(f"scrape_quiz.py version {VERSION}")
    store = Store(args.out, args.title, ignore_numbers=not args.keep_number_variants)
    print(f"Loaded {len(store.questions)} previously saved questions from {store.dir}")

    with sync_playwright() as pw:
        ctx = pw.chromium.launch_persistent_context(
            args.profile, headless=False, channel=args.browser,
            executable_path=args.executable,
            viewport={"width": 1400, "height": 900},
        )
        page = ctx.pages[0] if ctx.pages else ctx.new_page()
        url_file = store.dir / "quiz_url.txt"
        url = args.url or (url_file.read_text(encoding="utf-8").strip() if url_file.exists() else "")
        details_url = None
        if url:
            print(f"Opening {url}")
            try:
                page.goto(url, wait_until="domcontentloaded")
                if find_visible(page, lambda f: [f.get_by_role(r, name=START_RE) for r in ("button", "link")],
                                30) is not None:
                    details_url = page.url
            except PlaywrightError as e:
                print(f"  could not open it: {e}")
        if details_url is None:
            page, details_url = wait_for_user_on_quiz_page(ctx)
        url_file.write_text(details_url, encoding="utf-8")
        print(f"Using quiz page: {details_url}\n")

        dry = 0
        for attempt in range(1, args.max_attempts + 1):
            page = current_page(ctx, page)
            try:
                shown, new = run_attempt(page, details_url, store, attempt, args)
            except (RuntimeError, PlaywrightTimeout, PlaywrightError) as e:
                print(f"Attempt {attempt}: error - {str(e).splitlines()[0]}")
                page = current_page(ctx, page)
                try:
                    page.screenshot(path=str(store.dir / f"error_attempt_{attempt}.png"))
                    (store.dir / f"error_attempt_{attempt}.html").write_text(page.content(), encoding="utf-8")
                    print("  Saved a screenshot + HTML of the page to the output folder.")
                except PlaywrightError:
                    pass
                page = back_to_quiz(ctx, page, details_url)
                if page is None:
                    page, details_url = wait_for_user_on_quiz_page(ctx)
                    url_file.write_text(details_url, encoding="utf-8")
                continue
            store.save()
            dry = 0 if new else dry + 1
            print(f"Attempt {attempt}: {shown} questions shown, {new} new, "
                  f"{len(store.questions)} unique total")
            if dry >= args.stop_after:
                print(f"No new questions in {dry} attempts in a row - the bank looks complete.")
                break

        store.save()
        print(f"\nDone. {len(store.questions)} unique questions saved to {store.docx_path.resolve()}")
        ctx.close()


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        sys.exit("\nStopped. Everything collected so far is saved.")
